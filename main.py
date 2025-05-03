@app.route("/generate-docx", methods=["POST"])
def generate_docx():
    try:
        data = request.get_json()
        ticket_text = data.get("ticket_text", "")

        if not ticket_text:
            return jsonify({"error": "ticket_text is required"}), 400

        # Create UI-only DOCX
        from docx import Document
        import tempfile, os

        ui_doc = Document()
        ui_doc.add_heading("UI Checklist", level=1)
        ui_doc.add_paragraph(ticket_text)

        full_doc = Document()
        full_doc.add_heading("UI + Backend Troubleshooting", level=1)
        full_doc.add_paragraph(ticket_text)

        ui_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
        full_file = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
        ui_doc.save(ui_file.name)
        full_doc.save(full_file.name)

        return jsonify({
            "ui_doc_link": f"/download?file={os.path.basename(ui_file.name)}",
            "full_doc_link": f"/download?file={os.path.basename(full_file.name)}"
        })

    except Exception as e:
        return jsonify({"error": str(e)}), 500
